"""
build_workbook.py - Builds excel/retail_bi_kpi_workbook.xlsx from the
verified outputs in outputs/ (produced by prep_and_verify.py).

Sheets:
  KPI Summary   - headline KPIs from kpi_summary.json
  Category      - sales_by_category.csv
  Sub-Category  - sales_by_subcategory.csv
  Region        - region_profitability.csv
  Monthly Trend - monthly_trend.csv
  Discount Bands- discount_band_profit.csv

Run AFTER prep_and_verify.py:   python python/build_workbook.py
Deterministic for a given set of outputs (no timestamps are written
into cells; values come straight from the verified CSV/JSON files).
"""

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUTPUTS_DIR = ROOT / "outputs"
XLSX_PATH = ROOT / "excel" / "retail_bi_kpi_workbook.xlsx"

TITLE_FONT = Font(bold=True, size=14, color="1F3864")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY_FMT = '#,##0.00'
PCT_FMT = '0.00'


def write_table(ws, rows, start_row=1):
    """Write a list-of-lists table with a styled header row."""
    for r, row in enumerate(rows, start=start_row):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if r == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center")
    # widen columns to fit content
    for c, _ in enumerate(rows[0], start=1):
        width = max(len(str(row[c - 1])) for row in rows) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 12), 60)


def read_csv_rows(name):
    with open(OUTPUTS_DIR / name, newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        body = []
        for row in reader:
            body.append([_num(v) for v in row])
    return [header] + body


def _num(v):
    """Convert numeric strings to int/float so Excel treats them as numbers."""
    try:
        if "." in v:
            return float(v)
        return int(v)
    except (ValueError, TypeError):
        return v


def main() -> None:
    summary = json.loads((OUTPUTS_DIR / "kpi_summary.json").read_text())
    kpi = summary["headline_kpis"]

    wb = Workbook()

    # ---- KPI Summary sheet ----
    ws = wb.active
    ws.title = "KPI Summary"
    ws["A1"] = "Retail BI & Profitability Analysis - KPI Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Source: Tableau Sample Superstore (star schema in SQLite, verified by SQL)"
    rows = [
        ["KPI", "Value"],
        ["Total Sales (USD)", kpi["total_sales"]],
        ["Total Profit (USD)", kpi["total_profit"]],
        ["Profit Margin %", kpi["profit_margin_pct"]],
        ["Orders", kpi["order_count"]],
        ["Order Lines", kpi["order_lines"]],
        ["Customers", kpi["customer_count"]],
        ["Products", kpi["product_count"]],
        ["Units Sold", kpi["units_sold"]],
        ["Average Order Value (USD)", kpi["avg_order_value"]],
        ["Average Discount", kpi["avg_discount"]],
        ["First Order Date", kpi["first_order_date"]],
        ["Last Order Date", kpi["last_order_date"]],
        ["Best Category by Profit", summary["best_category_by_profit"]["category"]],
        ["Best Region by Profit", summary["best_region_by_profit"]["region"]],
        ["Top Product by Profit", summary["top_product_by_profit"]["product_name"]],
        ["Biggest Loss-Making Product", summary["biggest_loss_making_product"]["product_name"]],
    ]
    write_table(ws, rows, start_row=4)
    for r in range(5, 5 + len(rows) - 1):
        val = ws.cell(row=r, column=2)
        if isinstance(val.value, float):
            val.number_format = MONEY_FMT
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 55

    # ---- Data sheets straight from the verified query outputs ----
    sheets = [
        ("Category", "sales_by_category.csv"),
        ("Sub-Category", "sales_by_subcategory.csv"),
        ("Region", "region_profitability.csv"),
        ("Monthly Trend", "monthly_trend.csv"),
        ("Discount Bands", "discount_band_profit.csv"),
    ]
    for title, filename in sheets:
        ws = wb.create_sheet(title)
        write_table(ws, read_csv_rows(filename))

    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_PATH)
    print(f"[excel] workbook written to {XLSX_PATH} with sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
